#!/usr/bin/env python3
"""Build role-sync-demo/deck.pptx: the 6-slide TOI deck, generated onto the
2026 corporate template (real master/layouts, so Google Slides shows the
authentic branding). Mirrors deck.src.html slide-for-slide.

Usage: python3 build_pptx.py [path-to-template.pptx]
"""

import copy
import sys
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = pathlib.Path(__file__).parent
OUT = HERE.parent / "deck.pptx"
TEMPLATE = pathlib.Path(
    sys.argv[1]
    if len(sys.argv) > 1
    else "/home/andrewnguyen/workspace/redpanda/dump/2026 Template_ ADP rebrand (minimalist theme).pptx"
)

VIDEO_URL = "https://drive.google.com/file/d/14H8Yg9Om6Epx_a_HHOsvG51v5kKptwPA/view?usp=drive_link"

INK = RGBColor(0x12, 0x18, 0x27)
SLATE = RGBColor(0x69, 0x70, 0x84)
WASH = RGBColor(0xF3, 0xF5, 0xF9)
LINE = RGBColor(0xE3, 0xE7, 0xEF)
RED = RGBColor(0xE2, 0x40, 0x1B)
ORANGE = RGBColor(0xED, 0x83, 0x46)
PEACH = RGBColor(0xFA, 0xEA, 0xE5)
GREEN = RGBColor(0x18, 0x80, 0x38)
GREEN_TINT = RGBColor(0xE9, 0xF3, 0xEC)
GREEN_DARK = RGBColor(0x14, 0x53, 0x2D)
RED_DARK = RGBColor(0x8C, 0x1D, 0x05)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF, GROTESK, MONO = "Instrument Serif", "Space Grotesk", "Roboto Mono"

# deck.src.html is laid out on a 1280x720 grid; transplant its metrics.
EMU_PER_PX = 9144000 / 1280


def px(v):
    return Emu(round(v * EMU_PER_PX))


def runfmt(run, *, font=GROTESK, size=13, color=INK, bold=False, spacing=None):
    f = run.font
    f.name, f.size, f.bold = font, Pt(size), bold
    f.color.rgb = color
    if spacing is not None:  # tracking, in 1/100 pt (OOXML `spc`)
        f._rPr.set("spc", str(spacing))


def strip_style(sp):
    el = sp._element.find(qn("p:style"))
    if el is not None:
        sp._element.remove(el)


def para(tf, first=False, *, align=PP_ALIGN.LEFT, space_before=None, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    return p


def add_text(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return box, tf


def add_card(
    slide, x, y, w, h, *, fill, line_color=None, line_w=1.0, radius=0.10, dash=None
):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
    sp.adjustments[0] = radius
    strip_style(sp)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
        if dash:
            sp.line.dash_style = dash
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return sp, tf


def add_eyebrow(slide, text):
    _, tf = add_text(slide, 76, 30, 500, 26)
    r = para(tf, first=True).add_run()
    r.text = text.upper()
    runfmt(r, size=11, color=RED, bold=True, spacing=180)


def retitle(slide, text, *, top_px=62):
    title = slide.placeholders[0]
    title.left, title.top = px(76), px(top_px)
    title.width, title.height = px(1128), px(74)
    title.text_frame.text = text
    return title


def drop_unused_placeholders(slide, keep=(0,)):
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx not in keep and idx != 12:  # 12 = slide number, keep
            ph._element.getparent().remove(ph._element)


def rule_item(slide, x, y, w, head, why, *, head_size=13.5, why_size=11.5, bar_h=None):
    """The template's red-left-bar list item (HTML .rule / .bound / .how)."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, px(x), px(y + 2), px(3), px(bar_h or 58)
    )
    strip_style(bar)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()
    bar.shadow.inherit = False
    _, tf = add_text(slide, x + 16, y, w - 16, (bar_h or 58) + 8)
    p1 = para(tf, first=True, line=1.1)
    r = p1.add_run()
    r.text = head
    runfmt(r, size=head_size, bold=True)
    p2 = para(tf, space_before=3, line=1.15)
    r = p2.add_run()
    r.text = why
    runfmt(r, size=why_size, color=SLATE)
    return tf


def mono_row(slide, x, y, w, runs, *, missing=False):
    """A cluster-card row: mono text chip (HTML .row)."""
    sp, tf = add_card(
        slide,
        x,
        y,
        w,
        40,
        fill=None if missing else WHITE,
        line_color=ORANGE if missing else LINE,
        line_w=1.25 if missing else 1.0,
        radius=0.22,
        dash=MSO_LINE_DASH_STYLE.DASH if missing else None,
    )
    tf.margin_left, tf.margin_right = px(12), px(8)
    p = para(tf, first=True)
    for text, color, bold, font in runs:
        r = p.add_run()
        r.text = text
        runfmt(r, font=font, size=10.5, color=color, bold=bold)
    return sp


def main():
    prs = Presentation(TEMPLATE)
    layouts = {lo.name: lo for lo in prs.slide_layouts}
    lay_title = layouts["CUSTOM"]  # template slide 1 (hero art on layout)
    lay_content = layouts["ONE_COLUMN_TEXT_14"]  # template slide 3 (footer on layout)
    lay_divider = layouts["SECTION_HEADER_2_1"]  # template slide 2 (red bg on layout)

    # drop the template's 24 example slides; orphaned parts vanish on save
    sld_ids = prs.slides._sldIdLst
    for sld in list(sld_ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        sld_ids.remove(sld)

    # ---- 1 · TITLE ----------------------------------------------------
    s = prs.slides.add_slide(lay_title)
    s.placeholders[0].text_frame.text = "Shadowing Redpanda Roles"
    s.placeholders[1].width = px(700)
    tf = s.placeholders[1].text_frame
    tf.text = "July 2026"
    p = para(tf)
    p.add_run().text = "TOI — Shadow Linking Role Sync · Redpanda v26.2"
    s.placeholders[2].text_frame.text = "Andrew Nguyen — Core Engineering"

    # ---- 2 · THE GAP ---------------------------------------------------
    s = prs.slides.add_slide(lay_content)
    drop_unused_placeholders(s)
    add_eyebrow(s, "01 — Why")
    retitle(s, "Synced ACLs were inert without their roles")

    def cluster(x, head_b, head_rest, rows):
        add_card(s, x, 190, 470, 156, fill=WASH, line_color=LINE, radius=0.14)
        _, tf = add_text(s, x + 22, 204, 430, 20)
        p = para(tf, first=True)
        r = p.add_run()
        r.text = head_b.upper()
        runfmt(r, size=11, bold=True, spacing=140)
        r = p.add_run()
        r.text = f" · {head_rest.upper()}"
        runfmt(r, size=11, color=SLATE, spacing=140)
        for i, row in enumerate(rows):
            mono_row(
                s,
                x + 22,
                236 + i * 50,
                426,
                row["runs"],
                missing=row.get("missing", False),
            )

    acl_runs = [
        ("✓  ", GREEN, True, GROTESK),
        ("ACL: allow ", INK, False, MONO),
        ("RedpandaRole:analysts", RED, False, MONO),
    ]
    cluster(
        76,
        "Source",
        "production",
        [
            {"runs": acl_runs},
            {
                "runs": [
                    ("✓  ", GREEN, True, GROTESK),
                    ("Role: analysts {alice, bob}", INK, False, MONO),
                ]
            },
        ],
    )
    cluster(
        734,
        "Shadow",
        "DR",
        [
            {"runs": acl_runs},
            {
                "runs": [
                    ("✗  ", RED, True, GROTESK),
                    ("Role: analysts — missing", SLATE, False, MONO),
                ],
                "missing": True,
            },
        ],
    )

    arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, px(576), px(252), px(128), px(26))
    arrow.adjustments[0], arrow.adjustments[1] = 0.55, 0.55
    strip_style(arrow)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RED
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    for label, y in (("SHADOW LINK", 222), ("ACLS SYNCED — ROLES… NOT", 288)):
        _, tf = add_text(s, 526, y, 228, 30)
        p = para(tf, first=True, align=PP_ALIGN.CENTER, line=1.1)
        r = p.add_run()
        r.text = label
        runfmt(r, size=8.5, color=SLATE, bold=True, spacing=120)

    def verdict(x, tint, tag_color, body_color, tag, body):
        _, tf = add_card(s, x, 386, 552, 62, fill=tint, radius=0.18)
        tf.margin_left, tf.margin_right = px(20), px(10)
        p = para(tf, first=True)
        r = p.add_run()
        r.text = tag.upper() + "   "
        runfmt(r, size=10, color=tag_color, bold=True, spacing=110)
        r = p.add_run()
        r.text = body
        runfmt(r, size=12, color=body_color)

    verdict(
        76,
        GREEN_TINT,
        GREEN,
        GREEN_DARK,
        "Source",
        "alice produces to pageviews — allowed",
    )
    verdict(
        652, PEACH, RED, RED_DARK, "Shadow · after failover", "same request — denied"
    )

    # ---- 3 · WHAT ROLE SYNC DOES ----------------------------------------
    s = prs.slides.add_slide(lay_content)
    drop_unused_placeholders(s)
    add_eyebrow(s, "02 — What")
    retitle(s, "A full mirror, within a filter scope")

    verbs = [
        ("Create", "In scope on source, missing on shadow → created with members"),
        ("Update", "Membership differs → overwritten to match the source"),
        ("Delete", "On shadow but gone from source → removed"),
    ]
    for i, (verb, desc) in enumerate(verbs):
        x = 76 + i * 388
        add_card(s, x, 178, 352, 106, fill=WASH, line_color=LINE, radius=0.14)
        _, tf = add_text(s, x + 20, 192, 312, 82)
        p = para(tf, first=True)
        r = p.add_run()
        r.text = verb.upper()
        runfmt(r, size=12, color=RED, bold=True, spacing=140)
        p = para(tf, space_before=4, line=1.15)
        r = p.add_run()
        r.text = desc
        runfmt(r, size=10.5, color=INK)

    ax, ay, aw = 76, 316, 620
    add_card(s, ax, ay, aw, 182, fill=WASH, line_color=LINE, radius=0.09)
    _, tf = add_text(s, ax + 24, ay + 18, aw - 48, 150)
    p = para(tf, first=True)
    r = p.add_run()
    r.text = "In scope, the source is authoritative"
    runfmt(r, size=14, color=RED, bold=True)
    p = para(tf, space_before=7, line=1.25)
    for text, bold in (
        ("Add a member directly on the shadow → ", False),
        ("stripped", True),
        (" on the next sync. Delete an in-scope role on the shadow → ", False),
        ("recreated", True),
        (".", False),
    ):
        r = p.add_run()
        r.text = text
        runfmt(r, size=11.5, bold=bold)
    p = para(tf, space_before=7, line=1.25)
    r = p.add_run()
    r.text = "Roles outside the filter are never read, created, modified, or deleted. The filter is your blast-radius control."
    runfmt(r, size=11, color=SLATE)

    bounds = [
        (
            "Redpanda → Redpanda only",
            "Non-Redpanda source: the task parks; the rest of the link is unaffected.",
        ),
        (
            "Credentials aren't synced — by design",
            "Pre-provision identities on the shadow; role sync mirrors membership.",
        ),
        (
            "Opt-in",
            "Empty filters → nothing syncs. The task stays ACTIVE and does nothing.",
        ),
    ]
    for i, (head, why) in enumerate(bounds):
        rule_item(
            s, 730, 316 + i * 80, 474, head, why, head_size=11.5, why_size=10, bar_h=62
        )

    # ---- 4 · HOW IT READS ROLES -----------------------------------------
    s = prs.slides.add_slide(lay_content)
    drop_unused_placeholders(s)
    add_eyebrow(s, "03 — How")
    retitle(s, "One new Kafka API, in a reserved range")

    def wirebox(x, head, sub):
        add_card(s, x, 184, 300, 84, fill=WASH, line_color=LINE, radius=0.14)
        _, tf = add_text(s, x, 200, 300, 56, anchor=MSO_ANCHOR.TOP)
        p = para(tf, first=True, align=PP_ALIGN.CENTER)
        r = p.add_run()
        r.text = head
        runfmt(r, size=14, bold=True)
        p = para(tf, space_before=3, align=PP_ALIGN.CENTER)
        r = p.add_run()
        r.text = sub
        runfmt(r, size=10.5, color=SLATE)

    wirebox(76, "Shadow cluster", "roles migrator task")
    wirebox(904, "Source cluster", "Redpanda 26.2+")
    arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, px(408), px(216), px(464), px(20))
    arrow.adjustments[0], arrow.adjustments[1] = 0.5, 0.35
    strip_style(arrow)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RED
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    _, tf = add_text(s, 388, 186, 504, 22)
    p = para(tf, first=True, align=PP_ALIGN.CENTER)
    r = p.add_run()
    r.text = "DescribeRedpandaRoles · key 15000"
    runfmt(r, font=MONO, size=12, bold=True)
    _, tf = add_text(s, 388, 244, 504, 20)
    p = para(tf, first=True, align=PP_ALIGN.CENTER)
    r = p.add_run()
    r.text = "over the source's Kafka listener"
    runfmt(r, size=10.5, color=SLATE)

    hows = [
        (
            "Nothing standard to read",
            "Roles aren't in the Kafka data model — its security surface is ACLs, SCRAM, tokens, quotas.",
        ),
        (
            "Why a Kafka API",
            "The Admin API isn't reachable cross-cluster and isn't exposed in BYOC. The Kafka listener is the surface the link already uses.",
        ),
        (
            "Zero new permissions",
            "Authorized by cluster DESCRIBE — the same gate as DescribeAcls, which the link principal already holds.",
        ),
    ]
    for i, (head, why) in enumerate(hows):
        rule_item(
            s, 76 + i * 388, 300, 352, head, why, head_size=12, why_size=10, bar_h=92
        )

    _, tf = add_card(s, 76, 436, 1128, 72, fill=PEACH, radius=0.13)
    tf.margin_left, tf.margin_right = px(22), px(22)
    p = para(tf, first=True, line=1.15)
    r = p.add_run()
    r.text = "Consequence:  "
    runfmt(r, size=11.5, color=RED, bold=True)
    r = p.add_run()
    r.text = (
        "Apache Kafka / Confluent sources can't answer key 15000 → the role-sync task parks "
        "LINK_UNAVAILABLE; topic, ACL, and schema registry sync continue untouched. "
        "Role shadowing is Redpanda↔Redpanda DR."
    )
    runfmt(r, size=11.5, color=RED_DARK)

    # ---- 5 · CONFIG ------------------------------------------------------
    s = prs.slides.add_slide(lay_content)
    drop_unused_placeholders(s)
    add_eyebrow(s, "04 — Opting in")
    retitle(s, "Filters are the blast-radius control")

    add_card(s, 76, 184, 600, 296, fill=WASH, line_color=LINE, radius=0.08)
    _, tf = add_text(s, 106, 210, 546, 250)
    code = [
        [("name: ", INK), ("prod-shadow-link", GREEN)],
        [("# …topic / consumer / security options…", SLATE)],
        [("role_sync_options:", INK)],
        [("  interval: ", INK), ("30s", GREEN)],
        [("  role_name_filters:", INK)],
        [("    - pattern_type: ", INK), ("LITERAL", GREEN)],
        [("      filter_type: ", INK), ("INCLUDE", GREEN)],
        [("      name: ", INK), ("'*'", GREEN), ("   # wildcard: every role", SLATE)],
    ]
    for i, lruns in enumerate(code):
        p = para(tf, first=(i == 0), line=1.3)
        for text, color in lruns:
            r = p.add_run()
            r.text = text
            runfmt(r, font=MONO, size=11.5, color=color)

    rules = [
        (
            "Empty filters → nothing syncs.",
            "Role sync is opt-in: at least one INCLUDE.",
        ),
        ("LITERAL '*' is the only wildcard.", "Under PREFIX, * is just a character."),
        (
            "In scope, the source is authoritative.",
            "Created, updated — and deleted — to match the source.",
        ),
    ]
    for i, (head, why) in enumerate(rules):
        rule_item(s, 730, 196 + i * 92, 474, head, why, bar_h=64)

    # ---- 6 · DEMO DIVIDER ------------------------------------------------
    s = prs.slides.add_slide(lay_divider)
    s.placeholders[0].text_frame.text = "Demo"
    s.placeholders[
        1
    ].text_frame.text = "Two clusters, one link — create · converge · delete"

    btn, tf = add_card(
        s, 410, 546, 460, 54, fill=None, line_color=WHITE, line_w=1.5, radius=0.5
    )
    tf.word_wrap = False
    for m in ("margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.margin_left = px(145)
    p = para(tf, first=True, align=PP_ALIGN.LEFT)
    r = p.add_run()
    r.text = "Watch the demo recording"
    runfmt(r, size=12, color=WHITE, bold=True)
    btn.click_action.hyperlink.address = VIDEO_URL
    tri = s.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE, px(528), px(566), px(13), px(14)
    )
    strip_style(tri)
    tri.rotation = 90
    tri.fill.solid()
    tri.fill.fore_color.rgb = WHITE
    tri.line.fill.background()
    tri.shadow.inherit = False
    tri.click_action.hyperlink.address = VIDEO_URL

    prs.save(OUT)
    print(
        f"built {OUT} ({OUT.stat().st_size / 1024:.0f} KiB, {len(prs.slides._sldIdLst)} slides)"
    )


if __name__ == "__main__":
    main()
